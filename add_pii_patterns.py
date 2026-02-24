#!/usr/bin/env python3
"""
Script to add PII detection patterns to secret_rules.xlsx
"""

from openpyxl import load_workbook

# PII patterns to add
pii_patterns = [
    {
        'secret_type': 'Email Address',
        'pattern': r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b',
        'description': 'Email address (PII)',
        'confidence_level': 'HIGH'
    },
    {
        'secret_type': 'US Social Security Number',
        'pattern': r'\b\d{3}-\d{2}-\d{4}\b',
        'description': 'US SSN format XXX-XX-XXXX (PII)',
        'confidence_level': 'HIGH'
    },
    {
        'secret_type': 'Credit Card Number',
        'pattern': r'\b(?:\d{4}[-\s]?){3}\d{4}\b',
        'description': 'Credit card number (PII)',
        'confidence_level': 'HIGH'
    },
    {
        'secret_type': 'US Phone Number',
        'pattern': r'\b(?:\+?1[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b',
        'description': 'US phone number (PII)',
        'confidence_level': 'HIGH'
    },
    {
        'secret_type': 'IP Address',
        'pattern': r'\b(?:\d{1,3}\.){3}\d{1,3}\b',
        'description': 'IPv4 address (PII)',
        'confidence_level': 'LOW'
    },
    {
        'secret_type': 'US Passport Number',
        'pattern': r'\b[A-Z]{1,2}\d{6,9}\b',
        'description': 'US passport number (PII)',
        'confidence_level': 'HIGH'
    },
]

# Load the workbook
wb = load_workbook('secret_rules.xlsx')
ws = wb.active

# Find the last row
last_row = ws.max_row

# Add PII patterns
print(f"Adding {len(pii_patterns)} PII detection patterns...")
for i, pattern in enumerate(pii_patterns, start=1):
    row = last_row + i
    ws[f'A{row}'] = pattern['secret_type']
    ws[f'B{row}'] = pattern['pattern']
    ws[f'C{row}'] = pattern['description']
    ws[f'D{row}'] = pattern['confidence_level']
    print(f"  Added: {pattern['secret_type']}")

# Save the workbook
wb.save('secret_rules.xlsx')
print(f"\n✅ Successfully added {len(pii_patterns)} PII patterns to secret_rules.xlsx")
print(f"Total rules: {last_row + len(pii_patterns) - 1}")
